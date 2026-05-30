import logging
import os
from typing import Optional
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

logger = logging.getLogger(__name__)

class DocumentServer:
    def __init__(self, output_dir: str = "./output"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)
    
    def generate_word(self, content: str, title: str = "教案文档") -> str:
        try:
            doc = Document()
            
            title_para = doc.add_heading(title, 0)
            title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            lines = content.split('\n')
            for line in lines:
                line = line.strip()
                if not line:
                    continue
                
                if line.startswith('#'):
                    level = len(line) - len(line.lstrip('#'))
                    heading = line.lstrip('#').strip()
                    doc.add_heading(heading, level=min(level, 9))
                else:
                    para = doc.add_paragraph(line)
                    para.paragraph_format.line_spacing = 1.5
            
            filename = f"{title.replace(' ', '_')}.docx"
            filepath = os.path.join(self.output_dir, filename)
            doc.save(filepath)
            
            logger.info(f"Word 文档生成成功: {filepath}")
            return filepath
        except Exception as e:
            logger.error(f"生成 Word 文档失败: {e}")
            return ""
    
    def generate_markdown(self, content: str, title: str = "教案文档") -> str:
        try:
            md_content = f"# {title}\n\n{content}"
            filename = f"{title.replace(' ', '_')}.md"
            filepath = os.path.join(self.output_dir, filename)
            
            with open(filepath, 'w', encoding='utf-8') as f:
                f.write(md_content)
            
            logger.info(f"Markdown 文档生成成功: {filepath}")
            return filepath
        except Exception as e:
            logger.error(f"生成 Markdown 文档失败: {e}")
            return ""
    
    def save_to_file(self, content: str, filename: str, format: str = "md") -> str:
        if format == "docx":
            return self.generate_word(content, filename)
        else:
            return self.generate_markdown(content, filename)
